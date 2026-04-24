"""docx parser."""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph

from hks.errors import ExitCode, KSError
from hks.ingest.fingerprint import ParserFlags
from hks.ingest.models import ParsedDocument
from hks.ingest.office_common import (
    Segment,
    SkippedSegment,
    build_placeholder,
    merge_skipped,
    to_markdown_table,
)

_HEADING_RE = re.compile(r"heading\s*(\d+)", re.IGNORECASE)


def _clean_text(text: str) -> str:
    return " ".join(part.strip() for part in text.splitlines() if part.strip()).strip()


def _iter_blocks(document: DocumentObject) -> list[Paragraph | Table]:
    blocks: list[Paragraph | Table] = []
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            blocks.append(Paragraph(child, document))
        elif isinstance(child, CT_Tbl):
            blocks.append(Table(child, document))
    return blocks


def _extract_image_placeholders(paragraph: Paragraph) -> list[str]:
    payloads: list[str] = []
    for node in paragraph._element.xpath(".//*[local-name()='docPr']"):
        payload = str(node.get("descr") or node.get("title") or node.get("name") or "").strip()
        payloads.append(build_placeholder("image", payload))
    return payloads


def _extract_embedded_object_placeholders(paragraph: Paragraph) -> list[str]:
    payloads: list[str] = []
    for node in paragraph._element.xpath(".//*[local-name()='oleObject']"):
        payload = str(node.get("ProgID") or node.get("Type") or "").strip()
        payloads.append(build_placeholder("embedded_object", payload))
    return payloads


def _extract_smartart_placeholders(paragraph: Paragraph) -> list[str]:
    placeholders: list[str] = []
    for _node in paragraph._element.xpath(
        ".//*[local-name()='graphicData' and contains(@uri, 'diagram')]"
    ):
        placeholders.append(build_placeholder("smartart", ""))
    return placeholders


def _parse_paragraph(
    paragraph: Paragraph,
    *,
    title_ref: list[str],
    skipped_segments: list[SkippedSegment],
) -> list[Segment]:
    text = _clean_text(paragraph.text)
    style_name = paragraph.style.name if paragraph.style is not None else ""
    style_lower = style_name.lower()
    segments: list[Segment] = []

    heading_match = _HEADING_RE.match(style_name)
    if text:
        if heading_match:
            level = int(heading_match.group(1))
            if not title_ref[0]:
                title_ref[0] = text
            segments.append(
                Segment(
                    kind="heading",
                    text=f"{'#' * min(level + 1, 6)} {text}",
                    metadata={"level": level},
                )
            )
        elif (
            paragraph._p.pPr is not None
            and paragraph._p.pPr.numPr is not None
            or "list" in style_lower
        ):
            segments.append(Segment(kind="list_item", text=f"- {text}"))
        else:
            segments.append(Segment(kind="paragraph", text=text))

    placeholder_texts = (
        _extract_image_placeholders(paragraph)
        + _extract_embedded_object_placeholders(paragraph)
        + _extract_smartart_placeholders(paragraph)
    )
    for placeholder_text in placeholder_texts:
        if placeholder_text.startswith("[image: "):
            merge_skipped(skipped_segments, SkippedSegment(type="image"))
        elif placeholder_text.startswith("[embedded object: "):
            merge_skipped(skipped_segments, SkippedSegment(type="embedded_object"))
        else:
            merge_skipped(skipped_segments, SkippedSegment(type="smartart"))
        segments.append(Segment(kind="placeholder", text=placeholder_text))
    return segments


def _parse_table(table: Table) -> list[Segment]:
    rows = [[_clean_text(cell.text) for cell in row.cells] for row in table.rows]
    if not rows:
        return []
    header = [cell or f"col_{index + 1}" for index, cell in enumerate(rows[0])]
    body_rows = [row for row in rows[1:] if any(cell for cell in row)]
    if not body_rows:
        return [Segment(kind="table_row", text=to_markdown_table(header, []))]
    return [
        Segment(
            kind="table_row",
            text=to_markdown_table(header, [row]),
            metadata={"header_row": header},
        )
        for row in body_rows
    ]


def _append_macro_placeholder(
    path: Path,
    segments: list[Segment],
    skipped_segments: list[SkippedSegment],
) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            if any(name.endswith("vbaProject.bin") for name in archive.namelist()):
                merge_skipped(skipped_segments, SkippedSegment(type="macros"))
                segments.append(Segment(kind="placeholder", text=build_placeholder("macros")))
    except (OSError, zipfile.BadZipFile):  # pragma: no cover - handled by main parse path
        return


def parse(path: Path, flags: ParserFlags) -> ParsedDocument:
    del flags
    try:
        document = Document(str(path))
    except (PackageNotFoundError, ValueError, zipfile.BadZipFile) as exc:
        raise KSError(
            f"corrupt file: {path}",
            exit_code=ExitCode.DATAERR,
            code="CORRUPT",
            details=[str(exc)],
        ) from exc

    title_ref = [""]
    skipped_segments: list[SkippedSegment] = []
    segments: list[Segment] = []
    for block in _iter_blocks(document):
        if isinstance(block, Paragraph):
            segments.extend(
                _parse_paragraph(block, title_ref=title_ref, skipped_segments=skipped_segments)
            )
        else:
            segments.extend(_parse_table(block))
    _append_macro_placeholder(path, segments, skipped_segments)
    title = title_ref[0] or path.stem.replace("-", " ").replace("_", " ").strip() or path.stem
    return ParsedDocument(
        title=title,
        body="",
        format="docx",
        segments=segments,
        skipped_segments=skipped_segments,
    )
